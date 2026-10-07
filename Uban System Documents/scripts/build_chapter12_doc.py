# -*- coding: utf-8 -*-
r"""
Generate Chapter 12 Document (Word .docx and Markdown .md)
Strictly adheres to:
1. Heading 1: 16 pt, Heading 2: 14 pt, Body: 12 pt.
2. Fonts: Chinese DFKai-SB (標楷體), English Times New Roman.
3. Typography: All 12pt body text and table cells strictly EXACTLY 18pt line spacing, space before 0pt, space after 0pt.
4. Tables: All text strictly 12 pt, vertical alignment center.
   - Attribute column: centered horizontally, background #EBF1F5.
   - Content column: left-aligned.
   - cantSplit on every row, tblHeader on header rows.
   - Captions and Figure captions: bold=False, keep_with_next=True.
5. Images: Authentic screenshots from Flutter Native Skia pipeline, centered at 3.2 inches.
6. Section breaks: Each subsection (12-1, 12-2, 12-3) separated by Section Break (Next Page).
7. Official System Name: "Uban". 0 occurrences of "Uban 居家高齡長者智慧守護系統", "高熵", "收件匣".
"""

import os
import shutil
import docx
from docx.shared import Pt, RGBColor, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.enum.section import WD_SECTION_START
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

from chapter12_data import (
    IMAGES_DIR,
    SECTION_12_1_INTRO,
    MANUAL_ITEMS_12_1,
    SECTION_12_2_INTRO,
    MANUAL_ITEMS_12_2,
    SECTION_12_3_INTRO,
    MANUAL_ITEMS_12_3,
)

def set_cell_background(cell, color_hex):
    shading_xml = f'<w:shd {nsdecls("w")} w:fill="{color_hex}"/>'
    cell._tc.get_or_add_tcPr().append(parse_xml(shading_xml))

def set_cell_margins(cell, top=60, bottom=60, left=90, right=90):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = parse_xml(
        f'<w:tcMar {nsdecls("w")}>'
        f'<w:top w:w="{top}" w:type="dxa"/>'
        f'<w:bottom w:w="{bottom}" w:type="dxa"/>'
        f'<w:left w:w="{left}" w:type="dxa"/>'
        f'<w:right w:w="{right}" w:type="dxa"/>'
        f'</w:tcMar>'
    )
    tcPr.append(tcMar)

def set_table_borders(table, color="B0C4DE", sz="4", val="single"):
    tblPr = table._tbl.tblPr
    borders_xml = f"""
    <w:tblBorders {nsdecls("w")}>
        <w:top w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>
        <w:left w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>
        <w:bottom w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>
        <w:right w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>
        <w:insideH w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>
        <w:insideV w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>
    </w:tblBorders>
    """
    tblPr.append(parse_xml(borders_xml))

def format_run(run, text, font_east="標楷體", font_ascii="Times New Roman", size_pt=12, bold=False, color_rgb=(0, 0, 0)):
    run.text = text
    run.font.name = font_ascii
    run.font.size = Pt(size_pt)
    run.bold = bold
    run.font.color.rgb = RGBColor(*color_rgb)
    rPr = run._r.get_or_add_rPr()
    rFonts = rPr.find(qn('w:rFonts'))
    if rFonts is None:
        rFonts = OxmlElement('w:rFonts')
        rPr.append(rFonts)
    rFonts.set(qn('w:eastAsia'), font_east)
    rFonts.set(qn('w:ascii'), font_ascii)
    rFonts.set(qn('w:hAnsi'), font_ascii)

def add_heading_1(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(16)
    p.paragraph_format.space_after = Pt(8)
    p.paragraph_format.keep_with_next = True
    r = p.add_run()
    format_run(r, text, font_east="標楷體", font_ascii="Times New Roman", size_pt=16, bold=True, color_rgb=(0, 51, 102))
    return p

def add_heading_2(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(14)
    p.paragraph_format.space_after = Pt(6)
    p.paragraph_format.keep_with_next = True
    r = p.add_run()
    format_run(r, text, font_east="標楷體", font_ascii="Times New Roman", size_pt=14, bold=True, color_rgb=(20, 60, 110))
    return p

def add_paragraph(doc, text, bold_prefix="", indent=True):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.EXACTLY
    p.paragraph_format.line_spacing = Pt(18)
    if indent:
        p.paragraph_format.first_line_indent = Pt(24)
    if bold_prefix:
        r_pre = p.add_run()
        format_run(r_pre, bold_prefix, font_east="標楷體", font_ascii="Times New Roman", size_pt=12, bold=True, color_rgb=(0, 0, 0))
    r_body = p.add_run()
    format_run(r_body, text, font_east="標楷體", font_ascii="Times New Roman", size_pt=12, bold=False, color_rgb=(30, 30, 30))
    return p

def add_section_break(doc):
    new_sect = doc.add_section(WD_SECTION_START.NEW_PAGE)
    new_sect.top_margin = Inches(1.0)
    new_sect.bottom_margin = Inches(1.0)
    new_sect.left_margin = Inches(1.0)
    new_sect.right_margin = Inches(1.0)
    return new_sect

def add_caption(doc, text, is_table=True):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    if is_table:
        p.paragraph_format.space_before = Pt(6)
        p.paragraph_format.space_after = Pt(2)
    else:
        p.paragraph_format.space_before = Pt(2)
        p.paragraph_format.space_after = Pt(8)
    p.paragraph_format.keep_with_next = True
    r = p.add_run()
    # 遵照需求 6：所有圖、表標題不使用粗體，並且應與表格/圖置於同一頁
    format_run(r, text, font_east="標楷體", font_ascii="Times New Roman", size_pt=12, bold=False, color_rgb=(20, 40, 80))

def add_manual_item(doc, item):
    # 1. Table caption
    add_caption(doc, f"{item['table_id']} {item['title']}", is_table=True)

    # 2. Table creation
    rows_def = [
        ("項目", "系統規格與操作詳細說明"),
        ("畫面名稱", item["screen_name"]),
        ("適用角色", item["target_role"]),
        ("核心功能", item["core_functions"]),
        ("操作步驟", item["operation_steps"]),
    ]

    table = doc.add_table(rows=len(rows_def), cols=2)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    set_table_borders(table, color="B0C4DE", sz="4", val="single")

    col_widths = [Inches(1.5), Inches(5.0)]

    # 表頭設定
    header_tr = table.rows[0]._tr.get_or_add_trPr()
    header_tr.append(parse_xml(f'<w:tblHeader {nsdecls("w")}/>'))

    for row_idx, (label, content) in enumerate(rows_def):
        row = table.rows[row_idx]
        trPr = row._tr.get_or_add_trPr()
        trPr.append(parse_xml(f'<w:cantSplit {nsdecls("w")}/>'))

        is_header = (row_idx == 0)

        # Col 0: Label
        cell_0 = row.cells[0]
        cell_0.width = col_widths[0]
        cell_0.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
        set_cell_margins(cell_0, top=50, bottom=50, left=80, right=80)
        if is_header:
            set_cell_background(cell_0, "D9E1F2")
        else:
            set_cell_background(cell_0, "EBF1F5")

        p0 = cell_0.paragraphs[0]
        p0.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p0.paragraph_format.space_before = Pt(0)
        p0.paragraph_format.space_after = Pt(0)
        p0.paragraph_format.line_spacing_rule = WD_LINE_SPACING.EXACTLY
        p0.paragraph_format.line_spacing = Pt(18)
        r0 = p0.add_run()
        format_run(r0, label, font_east="標楷體", font_ascii="Times New Roman", size_pt=12, bold=True, color_rgb=(20, 40, 80))

        # Col 1: Content
        cell_1 = row.cells[1]
        cell_1.width = col_widths[1]
        cell_1.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
        set_cell_margins(cell_1, top=50, bottom=50, left=80, right=80)
        if is_header:
            set_cell_background(cell_1, "D9E1F2")
            p1 = cell_1.paragraphs[0]
            p1.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p1.paragraph_format.space_before = Pt(0)
            p1.paragraph_format.space_after = Pt(0)
            p1.paragraph_format.line_spacing_rule = WD_LINE_SPACING.EXACTLY
            p1.paragraph_format.line_spacing = Pt(18)
            r1 = p1.add_run()
            format_run(r1, content, font_east="標楷體", font_ascii="Times New Roman", size_pt=12, bold=True, color_rgb=(20, 40, 80))
        else:
            set_cell_background(cell_1, "FFFFFF")
            lines = content.split("\n")
            for line_idx, line in enumerate(lines):
                if line_idx == 0:
                    p1 = cell_1.paragraphs[0]
                else:
                    p1 = cell_1.add_paragraph()
                p1.alignment = WD_ALIGN_PARAGRAPH.LEFT
                p1.paragraph_format.space_before = Pt(0)
                p1.paragraph_format.space_after = Pt(0)
                p1.paragraph_format.line_spacing_rule = WD_LINE_SPACING.EXACTLY
                p1.paragraph_format.line_spacing = Pt(18)
                r1 = p1.add_run()
                format_run(r1, line, font_east="標楷體", font_ascii="Times New Roman", size_pt=12, bold=False, color_rgb=(20, 20, 20))

    # 3. Add Screenshot Image
    img_path = os.path.join(IMAGES_DIR, item["img_file"])
    if os.path.exists(img_path):
        p_img = doc.add_paragraph()
        p_img.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_img.paragraph_format.space_before = Pt(6)
        p_img.paragraph_format.space_after = Pt(2)
        p_img.paragraph_format.keep_with_next = True
        run_img = p_img.add_run()
        run_img.add_picture(img_path, width=Inches(3.2))

        # 4. Figure Caption
        add_caption(doc, f"{item['fig_id']} {item['fig_title']}", is_table=False)
    else:
        print(f"WARNING: Image not found: {img_path}")

def generate_docx(output_path):
    print(f"Generating DOCX -> {output_path}")
    doc = docx.Document()

    # Set margins
    for sec in doc.sections:
        sec.top_margin = Inches(1.0)
        sec.bottom_margin = Inches(1.0)
        sec.left_margin = Inches(1.0)
        sec.right_margin = Inches(1.0)

    # Big Chapter Heading
    add_heading_1(doc, "第十二章 使用手冊")
    add_paragraph(
        doc,
        "本手冊旨在提供 Uban 之完整端對端操作指南。"
        "本系統完全基於居家照護與家庭情感陪伴場景設計，全面消除繁瑣管理的複雜性，"
        "以簡約、溫馨、安全的互動架構連結出門在外的子女家屬與獨居在家的長輩。"
        "手冊內容依據角色權限與使用場景，劃分為「12-1 系統通用登入與身分確認」、「12-2 長者端使用手冊」"
        "以及「12-3 家屬守護端使用手冊」三大單元。全書所有畫面皆擷取自真實 Flutter 跨平台應用程式，"
        "並依循畫面功能、操作程序與實際截圖進行規範化規格描述。",
        indent=True
    )

    # ─── 12-1 系統通用登入與身分確認 ───
    add_heading_2(doc, "12-1 系統通用登入與身分確認")
    add_paragraph(doc, SECTION_12_1_INTRO, indent=True)
    for item in MANUAL_ITEMS_12_1:
        add_manual_item(doc, item)

    # ─── 12-2 長者端使用手冊 ───
    add_section_break(doc)
    add_heading_2(doc, "12-2 長者端使用手冊")
    add_paragraph(doc, SECTION_12_2_INTRO, indent=True)
    for item in MANUAL_ITEMS_12_2:
        add_manual_item(doc, item)

    # ─── 12-3 家屬守護端使用手冊 ───
    add_section_break(doc)
    add_heading_2(doc, "12-3 家屬守護端使用手冊")
    add_paragraph(doc, SECTION_12_3_INTRO, indent=True)
    for item in MANUAL_ITEMS_12_3:
        add_manual_item(doc, item)

    os.makedirs(os.path.dirname(output_path), exist_ok=True)
    doc.save(output_path)
    print(f"DOCX successfully generated: {output_path}")

def generate_markdown(output_path):
    print(f"Generating Markdown -> {output_path}")
    md = []
    md.append("# 第十二章 使用手冊\n")
    md.append(
        "本手冊旨在提供 Uban 之完整端對端操作指南。"
        "本系統完全基於居家照護與家庭情感陪伴場景設計，全面消除繁瑣管理的複雜性，"
        "以簡約、溫馨、安全的互動架構連結出門在外的子女家屬與獨居在家的長輩。"
        "手冊內容依據角色權限與使用場景，劃分為「12-1 系統通用登入與身分確認」、「12-2 長者端使用手冊」"
        "以及「12-3 家屬守護端使用手冊」三大單元。全書所有畫面皆擷取自真實 Flutter 跨平台應用程式，"
        "並依循畫面功能、操作程序與實際截圖進行規範化規格描述。\n"
    )

    def append_section(title, intro, items):
        md.append(f"## {title}\n")
        md.append(f"{intro}\n")
        for item in items:
            md.append(f"### {item['table_id']} {item['title']}\n")
            md.append("| 項目 | 系統規格與操作詳細說明 |")
            md.append("| :---: | :--- |")
            md.append(f"| **畫面名稱** | {item['screen_name']} |")
            md.append(f"| **適用角色** | {item['target_role']} |")
            core_fmt = item['core_functions'].replace('\n', '<br>')
            md.append(f"| **核心功能** | {core_fmt} |")
            step_fmt = item['operation_steps'].replace('\n', '<br>')
            md.append(f"| **操作步驟** | {step_fmt} |\n")
            md.append(f"![{item['fig_title']}](images/{item['img_file']})\n")
            md.append(f"<p align=\"center\">{item['fig_id']} {item['fig_title']}</p>\n")

    append_section("12-1 系統通用登入與身分確認", SECTION_12_1_INTRO, MANUAL_ITEMS_12_1)
    append_section("12-2 長者端使用手冊", SECTION_12_2_INTRO, MANUAL_ITEMS_12_2)
    append_section("12-3 家屬守護端使用手冊", SECTION_12_3_INTRO, MANUAL_ITEMS_12_3)

    os.makedirs(os.path.dirname(output_path), exist_ok=True)
    with open(output_path, "w", encoding="utf-8") as f:
        f.write("\n".join(md))
    print(f"Markdown successfully generated: {output_path}")

def main():
    base_dir = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
    target_dir = os.path.join(base_dir, "章節性文件", "Chapter 12")
    docx_path = os.path.join(target_dir, "Chapter 12.docx")
    md_path = os.path.join(target_dir, "Chapter 12.md")

    generate_docx(docx_path)
    generate_markdown(md_path)

    # Root copy
    root_docx = os.path.join(base_dir, "章節性文件", "Chapter 12.docx")
    root_md = os.path.join(base_dir, "章節性文件", "Chapter 12.md")

    shutil.copyfile(docx_path, root_docx)
    shutil.copyfile(md_path, root_md)
    print(f"Copied to root files:\n  {root_docx}\n  {root_md}")

if __name__ == "__main__":
    main()
