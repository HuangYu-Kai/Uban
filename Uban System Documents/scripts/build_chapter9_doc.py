# -*- coding: utf-8 -*-
r"""
Generate Chapter 9 Document (Word .docx and Markdown .md)
Strictly adheres to:
1. Heading 1: 16 pt, Heading 2: 14 pt, Body: 12 pt.
2. Fonts: Chinese DFKai-SB (標楷體), English Times New Roman.
3. Typography: All 12pt body text and table cells strictly EXACTLY 18pt line spacing, space before 0pt, space after 0pt.
4. Tables: All text 12 pt, vertical alignment center.
   - Non-detailed overview tables (9-1-1 ~ 9-1-3, 9-3-1 ~ 9-3-3, 9-4-1): horizontally centered.
   - Detailed code specification tables (9-2-1 ~ 9-2-8): left column centered label, right column left-aligned content.
   - Captions: bold=False, keep_with_next=True.
   - cantSplit on every row, tblHeader on header rows.
   - Table 9-1-3: page break before table to avoid split at page bottom.
   - Tables 9-2-1 ~ 9-2-8: page break before each code table to avoid split.
   - Table 9-3-3: compact margins (35 dxa) ensuring all 6 rows fit on one page.
5. Section breaks: Each subsection (9-1, 9-2, 9-3, 9-4) separated by Section Break (Next Page).
6. Official System Name: "Uban". 0 occurrences of "Uban 居家高齡長者智慧守護系統", "（逐行註解）", "高熵".
"""

import os
import shutil
import docx
from docx.shared import Pt, RGBColor, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.enum.section import WD_SECTION_START
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

from chapter9_data import (
    FRONTEND_PROGRAM_LIST,
    BACKEND_PROGRAM_LIST,
    ADMIN_PROGRAM_LIST,
    CODE_SNIPPETS_9_2,
    DEPENDENCIES_FLUTTER,
    DEPENDENCIES_BACKEND,
    DEPENDENCIES_ADMIN,
    EXTERNAL_API_LIST,
)

def set_cell_background(cell, color_hex):
    shading_xml = f'<w:shd {nsdecls("w")} w:fill="{color_hex}"/>'
    cell._tc.get_or_add_tcPr().append(parse_xml(shading_xml))

def set_cell_margins(cell, top=60, bottom=60, left=90, right=90):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = parse_xml(
        f'<w:tcMar {nsdecls("w")}>'
        f'<w:top w:w="{top}" w:type="dxa"/>'
        f'<w:bottom w:w="{bottom}" w:type="dxa"/>'
        f'<w:left w:w="{left}" w:type="dxa"/>'
        f'<w:right w:w="{right}" w:type="dxa"/>'
        f'</w:tcMar>'
    )
    tcPr.append(tcMar)

def set_table_borders(table, color="B0C4DE", sz="4", val="single"):
    tblPr = table._tbl.tblPr
    borders_xml = f"""
    <w:tblBorders {nsdecls("w")}>
        <w:top w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>
        <w:left w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>
        <w:bottom w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>
        <w:right w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>
        <w:insideH w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>
        <w:insideV w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>
    </w:tblBorders>
    """
    tblPr.append(parse_xml(borders_xml))

def format_run(run, text, font_east="標楷體", font_ascii="Times New Roman", size_pt=12, bold=False, color_rgb=(0, 0, 0)):
    run.text = text
    run.font.name = font_ascii
    run.font.size = Pt(size_pt)
    run.bold = bold
    run.font.color.rgb = RGBColor(*color_rgb)
    rPr = run._r.get_or_add_rPr()
    rFonts = rPr.find(qn('w:rFonts'))
    if rFonts is None:
        rFonts = OxmlElement('w:rFonts')
        rPr.append(rFonts)
    rFonts.set(qn('w:eastAsia'), font_east)
    rFonts.set(qn('w:ascii'), font_ascii)
    rFonts.set(qn('w:hAnsi'), font_ascii)

def add_heading_1(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(16)
    p.paragraph_format.space_after = Pt(8)
    p.paragraph_format.keep_with_next = True
    r = p.add_run()
    format_run(r, text, font_east="標楷體", font_ascii="Times New Roman", size_pt=16, bold=True, color_rgb=(0, 51, 102))
    return p

def add_heading_2(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(14)
    p.paragraph_format.space_after = Pt(6)
    p.paragraph_format.keep_with_next = True
    r = p.add_run()
    format_run(r, text, font_east="標楷體", font_ascii="Times New Roman", size_pt=14, bold=True, color_rgb=(20, 60, 110))
    return p

def add_paragraph(doc, text, bold_prefix="", indent=True):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.EXACTLY
    p.paragraph_format.line_spacing = Pt(18)
    if indent:
        p.paragraph_format.first_line_indent = Pt(24)
    if bold_prefix:
        r_pre = p.add_run()
        format_run(r_pre, bold_prefix, font_east="標楷體", font_ascii="Times New Roman", size_pt=12, bold=True, color_rgb=(0, 0, 0))
    r_body = p.add_run()
    format_run(r_body, text, font_east="標楷體", font_ascii="Times New Roman", size_pt=12, bold=False, color_rgb=(30, 30, 30))
    return p

def add_section_break(doc):
    new_sect = doc.add_section(WD_SECTION_START.NEW_PAGE)
    new_sect.top_margin = Inches(1.0)
    new_sect.bottom_margin = Inches(1.0)
    new_sect.left_margin = Inches(1.0)
    new_sect.right_margin = Inches(1.0)
    return new_sect

def create_centered_table(doc, table_id, title_text, headers, rows_data, col_widths=None):
    """非詳細說明表格：垂直置中、內容水平靠中對齊，字體 12 點"""
    p_caption = doc.add_paragraph()
    p_caption.paragraph_format.space_before = Pt(6)
    p_caption.paragraph_format.space_after = Pt(2)
    p_caption.paragraph_format.keep_with_next = True
    p_caption.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run_cap = p_caption.add_run()
    format_run(run_cap, f"表 {table_id} {title_text}", font_east="標楷體", font_ascii="Times New Roman", size_pt=12, bold=False, color_rgb=(20, 40, 80))

    table = doc.add_table(rows=len(rows_data) + 1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    set_table_borders(table, color="B0C4DE", sz="4", val="single")

    if col_widths and len(col_widths) == len(headers):
        for row in table.rows:
            for idx, width in enumerate(col_widths):
                row.cells[idx].width = width

    # 表頭設定
    header_tr = table.rows[0]._tr.get_or_add_trPr()
    header_tr.append(parse_xml(f'<w:tblHeader {nsdecls("w")}/>'))
    header_tr.append(parse_xml(f'<w:cantSplit {nsdecls("w")}/>'))

    for idx, h_text in enumerate(headers):
        cell = table.rows[0].cells[idx]
        set_cell_background(cell, "EBF1F5")
        if table_id == "9-3-3":
            set_cell_margins(cell, top=35, bottom=35, left=60, right=60)
        else:
            set_cell_margins(cell, top=60, bottom=60, left=90, right=90)
        cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(0)
        p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.EXACTLY
        p.paragraph_format.line_spacing = Pt(18)
        r = p.add_run()
        format_run(r, h_text, font_east="標楷體", font_ascii="Times New Roman", size_pt=12, bold=True, color_rgb=(20, 40, 80))

    # 資料列
    for row_idx, r_data in enumerate(rows_data):
        row = table.rows[row_idx + 1]
        trPr = row._tr.get_or_add_trPr()
        trPr.append(parse_xml(f'<w:cantSplit {nsdecls("w")}/>'))
        bg_color = "FAFAFA" if row_idx % 2 == 1 else "FFFFFF"
        for col_idx, val in enumerate(r_data):
            cell = row.cells[col_idx]
            set_cell_background(cell, bg_color)
            if table_id == "9-3-3":
                set_cell_margins(cell, top=35, bottom=35, left=60, right=60)
            else:
                set_cell_margins(cell, top=60, bottom=60, left=80, right=80)
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            p = cell.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.space_before = Pt(0)
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.EXACTLY
            p.paragraph_format.line_spacing = Pt(18)
            r = p.add_run()
            format_run(r, str(val), font_east="標楷體", font_ascii="Times New Roman", size_pt=12, bold=False, color_rgb=(40, 40, 40))

def create_code_spec_table(doc, item):
    """9-2 程式規格詳細說明表格：垂直置中、右欄靠左對齊，字體 12 點"""
    p_caption = doc.add_paragraph()
    p_caption.paragraph_format.space_before = Pt(6)
    p_caption.paragraph_format.space_after = Pt(2)
    p_caption.paragraph_format.keep_with_next = True
    p_caption.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run_cap = p_caption.add_run()
    format_run(run_cap, f"{item['table_id']} {item['title']}", font_east="標楷體", font_ascii="Times New Roman", size_pt=12, bold=False, color_rgb=(20, 40, 80))

    fields = [
        ("檔案來源", item["file_source"], False),
        ("設計目的", item["design_purpose"], False),
        ("核心程式碼", item["code_text"], True),
    ]

    table = doc.add_table(rows=len(fields), cols=2)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    set_table_borders(table, color="B0C4DE", sz="6", val="single")

    col_widths = [Inches(1.4), Inches(5.1)]
    for row in table.rows:
        trPr = row._tr.get_or_add_trPr()
        trPr.append(parse_xml(f'<w:cantSplit {nsdecls("w")}/>'))
        for idx, width in enumerate(col_widths):
            row.cells[idx].width = width

    for row_idx, (label, val, is_code) in enumerate(fields):
        cell_lbl = table.rows[row_idx].cells[0]
        cell_val = table.rows[row_idx].cells[1]

        set_cell_background(cell_lbl, "EBF1F5")
        set_cell_background(cell_val, "FFFFFF")
        set_cell_margins(cell_lbl, 60, 60, 90, 90)
        set_cell_margins(cell_val, 60, 60, 90, 90)

        cell_lbl.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
        cell_val.vertical_alignment = WD_ALIGN_VERTICAL.CENTER

        # 左欄標籤置中
        p_lbl = cell_lbl.paragraphs[0]
        p_lbl.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_lbl.paragraph_format.space_before = Pt(0)
        p_lbl.paragraph_format.space_after = Pt(0)
        p_lbl.paragraph_format.line_spacing_rule = WD_LINE_SPACING.EXACTLY
        p_lbl.paragraph_format.line_spacing = Pt(18)
        lines_lbl = label.split("\n")
        for idx, l in enumerate(lines_lbl):
            if idx > 0:
                p_lbl = cell_lbl.add_paragraph()
                p_lbl.alignment = WD_ALIGN_PARAGRAPH.CENTER
                p_lbl.paragraph_format.space_before = Pt(0)
                p_lbl.paragraph_format.space_after = Pt(0)
                p_lbl.paragraph_format.line_spacing_rule = WD_LINE_SPACING.EXACTLY
                p_lbl.paragraph_format.line_spacing = Pt(18)
            r_lbl = p_lbl.add_run()
            format_run(r_lbl, l, font_east="標楷體", font_ascii="Times New Roman", size_pt=12, bold=True, color_rgb=(20, 40, 80))

        # 右欄靠左
        p_val = cell_val.paragraphs[0]
        p_val.alignment = WD_ALIGN_PARAGRAPH.LEFT
        p_val.paragraph_format.space_before = Pt(0)
        p_val.paragraph_format.space_after = Pt(0)
        p_val.paragraph_format.line_spacing_rule = WD_LINE_SPACING.EXACTLY
        p_val.paragraph_format.line_spacing = Pt(18)

        if is_code:
            lines = val.split("\n")
            for l_idx, line in enumerate(lines):
                if l_idx > 0:
                    p_val = cell_val.add_paragraph()
                    p_val.alignment = WD_ALIGN_PARAGRAPH.LEFT
                    p_val.paragraph_format.space_before = Pt(0)
                    p_val.paragraph_format.space_after = Pt(0)
                    p_val.paragraph_format.line_spacing_rule = WD_LINE_SPACING.EXACTLY
                    p_val.paragraph_format.line_spacing = Pt(18)
                r_val = p_val.add_run()
                format_run(r_val, line, font_east="標楷體", font_ascii="Times New Roman", size_pt=12, bold=False, color_rgb=(20, 20, 20))
        else:
            r_val = p_val.add_run()
            format_run(r_val, val, font_east="標楷體", font_ascii="Times New Roman", size_pt=12, bold=False, color_rgb=(30, 30, 30))

def build_docx(target_docx_path):
    print(f"Building DOCX at {target_docx_path}...")
    doc = docx.Document()

    for s in doc.sections:
        s.top_margin = Inches(1.0)
        s.bottom_margin = Inches(1.0)
        s.left_margin = Inches(1.0)
        s.right_margin = Inches(1.0)

    # 第九章 軟體架構與程式規格（大標題 16 點）
    add_heading_1(doc, "第九章 軟體架構與程式規格")

    # 9-1 軟體架構與程式清單（小標題 14 點）
    add_heading_2(doc, "9-1 軟體架構與程式清單")
    add_paragraph(
        doc,
        "本系統「Uban」在軟體架構設計上採用現代化分層解耦模式，"
        "主要由四大子系統建構而成：(1) Flutter 跨平台雙端行動應用程式（長者生活端與家屬守護端）；"
        "(2) FastAPI 高併發非同步後端通訊伺服器；(3) React 18 開發者主控台；以及 (4) 邊緣與雲端智慧運算服務。"
        "各子系統之間透過標準 RESTful API、全雙工 Socket.IO 即時信令與點對點 WebRTC 串流進行安全通訊，"
        "徹底消除對繁瑣外部伺服器的相依性，提供高穩定、低延遲的家庭安全守護環境。"
    )

    # 表 9-1-1 前端程式清單
    create_centered_table(
        doc, "9-1-1", "Uban 行動端核心程式清單 (Flutter / Dart)",
        ["程式檔案名稱", "模組中文名稱", "核心功能與職責說明"],
        FRONTEND_PROGRAM_LIST,
        [Inches(2.2), Inches(1.6), Inches(2.7)]
    )

    # 表 9-1-2 後端程式清單
    create_centered_table(
        doc, "9-1-2", "Uban 後端核心服務程式清單 (FastAPI / Python)",
        ["程式檔案名稱", "服務中文名稱", "核心功能與職責說明"],
        BACKEND_PROGRAM_LIST,
        [Inches(1.8), Inches(1.8), Inches(2.9)]
    )

    # 表 9-1-3 開發者主控台清單（前加分頁符號，防止表格開頭 1~3 列被截斷至前頁）
    doc.add_page_break()
    create_centered_table(
        doc, "9-1-3", "Uban 開發者主控台程式清單 (React / TypeScript)",
        ["程式檔案名稱", "介面模組名稱", "核心功能與職責說明"],
        ADMIN_PROGRAM_LIST,
        [Inches(1.8), Inches(1.8), Inches(2.9)]
    )

    # 9-2 程式規格描述（小標題 14 點，使用下一頁分節符號）
    add_section_break(doc)
    add_heading_2(doc, "9-2 程式規格描述")
    add_paragraph(
        doc,
        "為具體闡明本系統核心模組之技術實作細節，本節精選系統中最具專業性與關鍵價值的 8 個現役程式碼片段。"
        "範疇涵蓋：長者端即時語音校正覆蓋層、家屬端戶外 GPS 移動軌跡回溯與停留點聚類、全雙工低延遲視訊信令連線、"
        "YOLOv8 骨架跌倒判定與 FCM 緊急推播、空間幾何多邊形射線室內區域定位、Google Gemini 陪伴對話意圖呼叫、"
        "長者戶外移動異常滯留安全警戒，以及長者與家屬端 4 位數專屬配對碼防衝突綁定機制。"
        "所有收錄之程式片段皆詳列檔案來源路徑、架構設計目的，並對每一行程式碼逐行標註繁體中文詳細解說。"
    )

    for item in CODE_SNIPPETS_9_2:
        doc.add_page_break()
        create_code_spec_table(doc, item)

    # 9-3 核心相依套件規格清單（小標題 14 點，使用下一頁分節符號）
    add_section_break(doc)
    add_heading_2(doc, "9-3 核心相依套件規格清單")
    add_paragraph(
        doc,
        "為確保系統在各運行環境具備最佳相容性與長期維護性，本專案嚴格遴選業界成熟且具備長期支援"
        "之高品質開源套件庫。以下分別整理行動端、後端核心服務與開發者主控台之核心相依套件規格清單："
    )

    # 表 9-3-1 行動端相依套件
    create_centered_table(
        doc, "9-3-1", "行動端核心相依套件規格清單 (Flutter / Dart)",
        ["套件名稱", "使用版本", "用途與系統職責說明"],
        DEPENDENCIES_FLUTTER,
        [Inches(1.8), Inches(1.3), Inches(3.4)]
    )

    # 表 9-3-2 後端相依套件
    create_centered_table(
        doc, "9-3-2", "後端核心服務相依套件規格清單 (FastAPI / Python)",
        ["套件名稱", "使用版本", "用途與系統職責說明"],
        DEPENDENCIES_BACKEND,
        [Inches(1.8), Inches(1.3), Inches(3.4)]
    )

    # 表 9-3-3 主控台相依套件（微調欄位 padding 35 dxa，確保 6 列完整同頁）
    create_centered_table(
        doc, "9-3-3", "開發者主控台相依套件規格清單 (React / TypeScript)",
        ["套件名稱", "使用版本", "用途與系統職責說明"],
        DEPENDENCIES_ADMIN,
        [Inches(1.8), Inches(1.3), Inches(3.4)]
    )

    # 9-4 外部整合 API 服務清單（小標題 14 點，使用下一頁分節符號）
    add_section_break(doc)
    add_heading_2(doc, "9-4 外部整合 API 服務清單")
    add_paragraph(
        doc,
        "Uban 系統整合多項具備高可用度之雲端與網路協定服務，為家庭成員提供全方位、不間斷的智慧守護。"
        "相關整合端點與通訊協定清單如表 9-4-1 所示："
    )

    create_centered_table(
        doc, "9-4-1", "外部整合 API 與雲端通訊協定服務清單",
        ["服務名稱", "通訊協定 / 格式", "整合目的與業務運作職責"],
        EXTERNAL_API_LIST,
        [Inches(2.2), Inches(1.4), Inches(2.9)]
    )

    os.makedirs(os.path.dirname(target_docx_path), exist_ok=True)
    doc.save(target_docx_path)
    print(f"[SUCCESS] Successfully generated docx: {target_docx_path}")

def build_markdown(target_md_path):
    print(f"Building Markdown at {target_md_path}...")
    lines = []
    lines.append("# 第九章 軟體架構與程式規格\n")

    # 9-1
    lines.append("## 9-1 軟體架構與程式清單\n")
    lines.append(
        "本系統「Uban」在軟體架構設計上採用現代化分層解耦模式，"
        "主要由四大子系統建構而成：(1) Flutter 跨平台雙端行動應用程式（長者生活端與家屬守護端）；"
        "(2) FastAPI 高併發非同步後端通訊伺服器；(3) React 18 開發者主控台；以及 (4) 邊緣與雲端智慧運算服務。"
        "各子系統之間透過標準 RESTful API、全雙工 Socket.IO 即時信令與點對點 WebRTC 串流進行安全通訊，"
        "徹底消除對繁瑣外部伺服器的相依性，提供高穩定、低延遲的家庭安全守護環境。\n"
    )

    # 表 9-1-1
    lines.append("### 表 9-1-1 Uban 行動端核心程式清單 (Flutter / Dart)\n")
    lines.append("| 程式檔案名稱 | 模組中文名稱 | 核心功能與職責說明 |")
    lines.append("| :---: | :---: | :---: |")
    for r in FRONTEND_PROGRAM_LIST:
        lines.append(f"| `{r[0]}` | {r[1]} | {r[2]} |")
    lines.append("\n")

    # 表 9-1-2
    lines.append("### 表 9-1-2 Uban 後端核心服務程式清單 (FastAPI / Python)\n")
    lines.append("| 程式檔案名稱 | 服務中文名稱 | 核心功能與職責說明 |")
    lines.append("| :---: | :---: | :---: |")
    for r in BACKEND_PROGRAM_LIST:
        lines.append(f"| `{r[0]}` | {r[1]} | {r[2]} |")
    lines.append("\n")

    # 表 9-1-3
    lines.append("### 表 9-1-3 Uban 開發者主控台程式清單 (React / TypeScript)\n")
    lines.append("| 程式檔案名稱 | 介面模組名稱 | 核心功能與職責說明 |")
    lines.append("| :---: | :---: | :---: |")
    for r in ADMIN_PROGRAM_LIST:
        lines.append(f"| `{r[0]}` | {r[1]} | {r[2]} |")
    lines.append("\n")

    # 9-2
    lines.append("## 9-2 程式規格描述\n")
    lines.append(
        "為具體闡明本系統核心模組之技術實作細節，本節精選系統中最具專業性與關鍵價值的 8 個現役程式碼片段。"
        "範疇涵蓋：長者端即時語音校正覆蓋層、家屬端戶外 GPS 移動軌跡回溯與停留點聚類、全雙工低延遲視訊信令連線、"
        "YOLOv8 骨架跌倒判定與 FCM 緊急推播、空間幾何多邊形射線室內區域定位、Google Gemini 陪伴對話意圖呼叫、"
        "長者戶外移動異常滯留安全警戒，以及長者與家屬端 4 位數專屬配對碼防衝突綁定機制。"
        "所有收錄之程式片段皆詳列檔案來源路徑、架構設計目的，並對每一行程式碼逐行標註繁體中文詳細解說。\n"
    )

    for item in CODE_SNIPPETS_9_2:
        lines.append(f"### {item['table_id']} {item['title']}\n")
        lines.append("| 項目 | 系統規格與技術詳細說明 |")
        lines.append("| :---: | :--- |")
        lines.append(f"| **檔案來源** | `{item['file_source']}` |")
        lines.append(f"| **設計目的** | {item['design_purpose']} |")
        code_md = item['code_text'].replace('\n', '<br>').replace(' ', '&nbsp;')
        lines.append(f"| **核心程式碼** | <code>{code_md}</code> |\n")

    # 9-3
    lines.append("## 9-3 核心相依套件規格清單\n")
    lines.append(
        "為確保系統在各運行環境具備最佳相容性與長期維護性，本專案嚴格遴選業界成熟且具備長期支援"
        "之高品質開源套件庫。以下分別整理行動端、後端核心服務與開發者主控台之核心相依套件規格清單：\n"
    )

    lines.append("### 表 9-3-1 行動端核心相依套件規格清單 (Flutter / Dart)\n")
    lines.append("| 套件名稱 | 使用版本 | 用途與系統職責說明 |")
    lines.append("| :---: | :---: | :---: |")
    for r in DEPENDENCIES_FLUTTER:
        lines.append(f"| `{r[0]}` | {r[1]} | {r[2]} |")
    lines.append("\n")

    lines.append("### 表 9-3-2 後端核心服務相依套件規格清單 (FastAPI / Python)\n")
    lines.append("| 套件名稱 | 使用版本 | 用途與系統職責說明 |")
    lines.append("| :---: | :---: | :---: |")
    for r in DEPENDENCIES_BACKEND:
        lines.append(f"| `{r[0]}` | {r[1]} | {r[2]} |")
    lines.append("\n")

    lines.append("### 表 9-3-3 開發者主控台相依套件規格清單 (React / TypeScript)\n")
    lines.append("| 套件名稱 | 使用版本 | 用途與系統職責說明 |")
    lines.append("| :---: | :---: | :---: |")
    for r in DEPENDENCIES_ADMIN:
        lines.append(f"| `{r[0]}` | {r[1]} | {r[2]} |")
    lines.append("\n")

    # 9-4
    lines.append("## 9-4 外部整合 API 服務清單\n")
    lines.append(
        "Uban 系統整合多項具備高可用度之雲端與網路協定服務，為家庭成員提供全方位、不間斷的智慧守護。"
        "相關整合端點與通訊協定清單如表 9-4-1 所示：\n"
    )

    lines.append("### 表 9-4-1 外部整合 API 與雲端通訊協定服務清單\n")
    lines.append("| 服務名稱 | 通訊協定 / 格式 | 整合目的與業務運作職責 |")
    lines.append("| :---: | :---: | :---: |")
    for r in EXTERNAL_API_LIST:
        lines.append(f"| {r[0]} | {r[1]} | {r[2]} |")
    lines.append("\n")

    os.makedirs(os.path.dirname(target_md_path), exist_ok=True)
    with open(target_md_path, "w", encoding="utf-8") as f:
        f.write("\n".join(lines))
    print(f"[SUCCESS] Successfully generated Markdown: {target_md_path}")

def main():
    base_dir = r"E:\114Project\Uban\Uban System Documents\章節性文件"
    sub_dir = os.path.join(base_dir, "Chapter 9")
    os.makedirs(sub_dir, exist_ok=True)

    target_docx = os.path.join(sub_dir, "Chapter 9.docx")
    target_md = os.path.join(sub_dir, "Chapter 9.md")

    build_docx(target_docx)
    build_markdown(target_md)

    root_docx = os.path.join(base_dir, "Chapter 9.docx")
    root_md = os.path.join(base_dir, "Chapter 9.md")
    try:
        shutil.copy2(target_docx, root_docx)
        print(f"[SUCCESS] Successfully copied to root docx: {root_docx}")
    except Exception as e:
        print(f"[WARNING] Could not copy to root docx: {e}")

    try:
        shutil.copy2(target_md, root_md)
        print(f"[SUCCESS] Successfully copied to root md: {root_md}")
    except Exception as e:
        print(f"[WARNING] Could not copy to root md: {e}")

if __name__ == "__main__":
    main()
